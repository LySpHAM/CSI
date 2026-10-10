from __future__ import annotations

import asyncio
import io
import os
import uuid
from datetime import UTC, datetime
from html import escape
from pathlib import Path

from sqlalchemy import select

from app.core.auth import load_principal
from app.core.config import get_settings
from app.core.resources import Resources
from app.core.storage import LocalStorage
from app.db.tenant import system_session, tenant_session
from app.domain.tabular import safe_cell
from app.models import ExportJob, Tenant
from app.models.enums import ExportKind, JobStatus
from app.schemas.feedback import FeedbackFilter
from app.services.analytics import AnalyticsService
from worker.celery_app import celery_app


async def _report_topic_data(service: AnalyticsService, filters: FeedbackFilter) -> dict:
    """Tổng hợp xu hướng, chất lượng theo chủ đề và ví dụ phản hồi tiêu cực."""
    quality = await service.topic_quality(filters)
    names = {row["topic_id"]: row["topic_name"] for row in quality}
    stats = {
        row["topic_id"]: {
            "total": row["total"],
            "positive": row["positive"],
            "negative": row["negative"],
            "urgent": row["urgent"],
            "csat": row["csat"],
            "recommendation_key": row["recommendation_key"],
            "sample": row["sample"],
        }
        for row in quality
    }

    trend_rows = await service.topics(filters, period="week")
    buckets = sorted({row["date"] for row in trend_rows})
    trend_counts: dict[str, dict[str, int]] = {}
    for row in trend_rows:
        trend_counts.setdefault(row["topic_id"], {})[row["date"]] = row["count"]

    issues = sorted(
        (
            (topic_id, data)
            for topic_id, data in stats.items()
            if data["negative"] >= 2
        ),
        key=lambda item: (item[1]["urgent"], item[1]["negative"], item[1]["total"]),
        reverse=True,
    )[:6]
    examples = {row["topic_id"]: row["sample"] for row in quality if row["sample"]}

    return {
        "names": names,
        "stats": stats,
        "buckets": buckets,
        "trends": trend_counts,
        "issues": issues,
        "examples": examples,
    }


async def render(
    service: AnalyticsService, filters: FeedbackFilter, kind: ExportKind, app_name: str
) -> bytes:
    summary = await service.overview(filters)
    topics = await service.topics(filters)
    output = io.BytesIO()
    if kind == ExportKind.XLSX:
        from openpyxl import Workbook

        workbook = Workbook(write_only=True)
        overview = workbook.create_sheet("Tóm tắt")
        overview.append([app_name])
        overview.append(["Từ", str(filters.start or ""), "Đến trước", str(filters.end or "")])
        overview.append(["Tổng phản hồi", summary["total_responses"]])
        overview.append(["CSAT", summary["csat"]])
        overview.append(["Khẩn cấp", summary["urgent_responses"]])
        for label, value in summary["sentiments"].items():
            overview.append(
                [
                    {"positive": "Tích cực", "negative": "Tiêu cực", "neutral": "Trung lập"}[label],
                    value["count"],
                    value["percent"],
                ]
            )
        detail = workbook.create_sheet("Chi tiết")
        detail.append(["Nội dung", "Sao", "Cảm xúc", "Chủ đề", "Khẩn cấp", "Nguồn", "Thời gian"])
        query = await service.analysis_scope(filters)
        result = await service.db.stream_scalars(query.execution_options(yield_per=200))
        async for row in result:
            detail.append(
                [
                    safe_cell(row.text),
                    float(row.rating) if row.rating is not None else None,
                    row.sentiment,
                    ",".join(str(t) for t in row.topic_ids),
                    row.is_urgent,
                    row.channel,
                    row.responded_at.isoformat(),
                ]
            )
        sheet = workbook.create_sheet("Theo chủ đề")
        sheet.append(["Kỳ", "Mã chủ đề", "Số nhận xét"])
        for topic in topics:
            sheet.append([topic["date"], topic["topic_id"], topic["count"]])
        workbook.save(output)
    else:
        from reportlab.graphics.charts.barcharts import VerticalBarChart
        from reportlab.graphics.shapes import Drawing
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, Table

        font = os.environ.get("REPORT_FONT_PATH", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
        if not await asyncio.to_thread(Path(font).is_file):
            raise RuntimeError("Thiếu font Unicode; đặt REPORT_FONT_PATH để xuất PDF tiếng Việt.")
        pdfmetrics.registerFont(TTFont("ReportUnicode", font))
        pdfmetrics.registerFontFamily(
            "ReportUnicode", normal="ReportUnicode", bold="ReportUnicode",
            italic="ReportUnicode", boldItalic="ReportUnicode",
        )
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = "ReportUnicode"
        ink = colors.HexColor("#17324d")
        green = colors.HexColor("#075b49")
        muted = colors.HexColor("#60758a")
        pale = colors.HexColor("#f4f8fb")
        body = ParagraphStyle(
            "ReportBody", parent=styles["BodyText"], fontName="ReportUnicode",
            fontSize=8.2, leading=11, textColor=ink,
        )
        small = ParagraphStyle(
            "ReportSmall", parent=body, fontSize=7.4, leading=9.5, textColor=muted,
        )
        section = ParagraphStyle(
            "ReportSection", parent=styles["Heading2"], fontName="ReportUnicode",
            fontSize=11, leading=14, textColor=green, spaceBefore=10, spaceAfter=6,
        )
        table_head = ParagraphStyle(
            "ReportTableHead", parent=body, fontName="ReportUnicode",
            fontSize=7.2, leading=9, textColor=colors.white,
        )
        topic_data = await _report_topic_data(service, filters)
        names, stats = topic_data["names"], topic_data["stats"]
        buckets, trends = topic_data["buckets"], topic_data["trends"]
        content = [
            Paragraph(escape(app_name), styles["Title"]),
            Paragraph("Báo cáo phản hồi khách hàng", styles["Heading1"]),
            Paragraph(
                escape(f"Từ {filters.start or 'đầu kỳ'} đến trước {filters.end or 'hiện tại'}"),
                body,
            ),
            Spacer(1, 10),
        ]
        summary_table = Table(
            [
                [Paragraph("Tổng phản hồi", body), str(summary["total_responses"])],
                [Paragraph("CSAT trung bình", body), f"{summary['csat']:.2f}/5" if summary["csat"] is not None else "—"],
                [Paragraph("Cảnh báo khẩn", body), str(summary["urgent_responses"])],
            ], colWidths=[150, 70],
        )
        summary_table.setStyle([
            ("FONTNAME", (0, 0), (-1, -1), "ReportUnicode"),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cfdae4")),
            ("BACKGROUND", (0, 0), (-1, -1), pale),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ])
        content.append(summary_table)
        chart = VerticalBarChart()
        chart.x = 40
        chart.y = 24
        chart.width = 350
        chart.height = 130
        chart.data = [
            [summary["sentiments"][s]["count"] for s in ("positive", "negative", "neutral")]
        ]
        chart.categoryAxis.categoryNames = ["Tích cực", "Tiêu cực", "Trung lập"]
        chart.categoryAxis.labels.fontName = "ReportUnicode"
        chart.bars[(0, 0)].fillColor = colors.green
        chart.bars[(0, 1)].fillColor = colors.red
        chart.bars[(0, 2)].fillColor = colors.grey
        drawing = Drawing(440, 175)
        drawing.add(chart)
        content.append(drawing)
        content.append(
            Paragraph(
                (
                    "Phân bố cảm xúc trên nhận xét văn bản đã phân tích. "
                    "Báo cáo không đưa dữ liệu định danh khách hàng vào."
                ),
                small,
            )
        )
        content.append(Paragraph("1. XU HƯỚNG CHỦ ĐỀ", section))
        trend_table = [[
            Paragraph("Chủ đề", table_head), Paragraph("Kỳ trước", table_head),
            Paragraph("Kỳ gần nhất", table_head), Paragraph("Thay đổi", table_head),
        ]]
        latest_bucket = buckets[-1] if buckets else None
        previous_bucket = buckets[-2] if len(buckets) > 1 else None
        trend_topics = sorted(
            (topic_id for topic_id in names if sum(trends.get(topic_id, {}).values()) > 0),
            key=lambda topic_id: sum(trends.get(topic_id, {}).values()),
            reverse=True,
        )[:10]
        for topic_id in trend_topics:
            counts = trends.get(topic_id, {})
            current = counts.get(latest_bucket, 0) if latest_bucket else 0
            previous = counts.get(previous_bucket, 0) if previous_bucket else 0
            if previous_bucket is None:
                change = "Chưa đủ kỳ"
            elif previous == 0 and current > 0:
                change = "Mới phát sinh"
            elif previous == 0:
                change = "Ổn định"
            else:
                delta = (current - previous) / previous * 100
                change = f"{'↑' if delta > 0 else '↓' if delta < 0 else '→'} {delta:+.0f}%"
            trend_table.append([
                Paragraph(escape(names[topic_id]), body),
                f"{previous} ({previous_bucket[:10]})" if previous_bucket else "—",
                f"{current} ({latest_bucket[:10]})" if latest_bucket else "—", change,
            ])
        if len(trend_table) == 1:
            trend_table.append([Paragraph("Chưa có dữ liệu chủ đề trong kỳ.", body), "", "", ""])
        topic_trend_table = LongTable(trend_table, colWidths=[190, 105, 105, 105], repeatRows=1)
        topic_trend_table.setStyle([
            ("BACKGROUND", (0, 0), (-1, 0), green),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, pale]),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d7e0e8")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTNAME", (0, 0), (-1, -1), "ReportUnicode"),
            ("FONTSIZE", (0, 1), (-1, -1), 7.5),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
        content.append(topic_trend_table)

        content.append(Paragraph("2. ĐÁNH GIÁ CHẤT LƯỢNG THEO TỪNG KHÍA CẠNH DỊCH VỤ", section))
        quality_table = [[
            Paragraph("Khía cạnh dịch vụ", table_head), Paragraph("Số phản hồi", table_head),
            Paragraph("Tích cực", table_head), Paragraph("Tiêu cực", table_head),
            Paragraph("Điểm CSAT", table_head), Paragraph("Đánh giá chung", table_head),
        ]]
        for topic_id in sorted(stats, key=lambda key: stats[key]["total"], reverse=True):
            item = stats[topic_id]
            total = item["total"]
            negative_rate = item["negative"] / total if total else 0
            positive_rate = item["positive"] / total if total else 0
            csat = item["csat"]
            if negative_rate >= 0.5 or (csat is not None and csat < 3):
                assessment = "Cảnh báo: cần ưu tiên"
            elif negative_rate >= 0.3 or (csat is not None and csat < 3.5):
                assessment = "Cần theo dõi"
            else:
                assessment = "Hoạt động ổn định"
            quality_table.append([
                Paragraph(escape(names[topic_id]), body), str(total),
                f"{positive_rate:.0%}", f"{negative_rate:.0%}",
                f"{csat:.1f}/5" if csat is not None else "—", Paragraph(assessment, small),
            ])
        if len(quality_table) == 1:
            quality_table.append([Paragraph("Chưa có dữ liệu phân tích theo chủ đề.", body), "", "", "", "", ""])
        quality = LongTable(quality_table, colWidths=[120, 58, 55, 55, 58, 159], repeatRows=1)
        quality.setStyle([
            ("BACKGROUND", (0, 0), (-1, 0), green),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, pale]),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#d7e0e8")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTNAME", (0, 0), (-1, -1), "ReportUnicode"),
            ("FONTSIZE", (0, 1), (-1, -1), 7.3),
            ("LEFTPADDING", (0, 0), (-1, -1), 5),
            ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ])
        content.append(quality)

        content.append(Paragraph("3. DANH SÁCH VẤN ĐỀ LẶP LẠI CẦN TẬP TRUNG KHẮC PHỤC", section))
        if topic_data["issues"]:
            negative_total = sum(item["negative"] for item in stats.values())
            for topic_id, item in topic_data["issues"]:
                share = item["negative"] / negative_total if negative_total else 0
                if item["urgent"] > 0 or (item["negative"] >= 3 and share >= 0.25):
                    priority, priority_color = "KHẨN CẤP", colors.HexColor("#c62828")
                elif share >= 0.15:
                    priority, priority_color = "ƯU TIÊN CAO", colors.HexColor("#d97706")
                else:
                    priority, priority_color = "THEO DÕI", muted
                details = (
                    f"<b>{escape(names[topic_id])}</b> &nbsp; "
                    f"<font color='#c2185b'>Lặp lại {item['negative']} lần ({share:.0%})</font>"
                )
                if topic_data["examples"].get(topic_id):
                    details += "<br/><font color='#60758a'><i>Trích dẫn mẫu: “" + escape(topic_data["examples"][topic_id]) + "”</i></font>"
                issue_card = Table(
                    [[Paragraph(details, body), Paragraph(priority, small)]],
                    colWidths=[440, 65],
                )
                issue_card.setStyle([
                    ("BACKGROUND", (0, 0), (-1, -1), pale),
                    ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#d7e0e8")),
                    ("TEXTCOLOR", (1, 0), (1, 0), priority_color),
                    ("ALIGN", (1, 0), (1, 0), "RIGHT"),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 8),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                    ("TOPPADDING", (0, 0), (-1, -1), 7),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                ])
                content.extend([issue_card, Spacer(1, 5)])
        else:
            content.append(Paragraph("Chưa ghi nhận chủ đề có từ hai phản hồi tiêu cực trở lên trong kỳ.", body))

        content.append(Paragraph("4. KHUYẾN NGHỊ CHIẾN LƯỢC CẢI THIỆN CHẤT LƯỢNG DỊCH VỤ", section))
        advice = {
            "payment": ("Kiểm tra luồng đặt hàng, voucher, số tiền cuối và đối chiếu giao dịch lỗi.", "Chuẩn hóa ưu đãi/thanh toán; theo dõi lỗi thanh toán, hóa đơn lệch và khiếu nại theo kênh."),
            "technical": ("Tái hiện lỗi trên đúng thiết bị/phiên bản; ghi thời điểm, thông báo lỗi và cách xử lý tạm.", "Theo dõi lỗi theo phiên bản, đặt hạn phản hồi sự cố và kiểm tra hồi quy trước khi phát hành."),
            "wait": ("Đo thời gian chờ theo giờ/ca, tìm điểm nghẽn và tăng hỗ trợ lúc cao điểm.", "Dự báo nhân sự; theo dõi thời gian chờ và phản hồi tiêu cực theo ca/chi nhánh."),
            "staff": ("Quản lý ca rà soát tình huống, trao đổi với nhân viên liên quan và thống nhất cách xử lý.", "Chuẩn hóa giao tiếp/bàn giao, huấn luyện theo tình huống và theo dõi khiếu nại lặp lại theo ca."),
            "price": ("Đối chiếu giá thấy được với hóa đơn; kiểm tra phụ phí, điều kiện giảm giá và đổi trả.", "Công khai thống nhất giá/phụ phí/điều kiện ưu đãi; theo dõi khiếu nại giá và giao dịch bỏ dở."),
            "product": ("Kiểm tra lô hàng/đơn được phản ánh, đóng gói và bảo quản; rà soát cùng mã sản phẩm.", "Tăng kiểm tra chất lượng tại khâu phát sinh; theo dõi lỗi, đổi trả theo mã hàng/lô."),
            "facility": ("Kiểm tra đúng khu vực/thời điểm; giao người phụ trách và hạn xử lý hạng mục.", "Lập lịch bảo trì/vệ sinh; theo dõi thời gian sửa và lỗi cơ sở vật chất tái diễn."),
            "other": ("Đối chiếu phản hồi mẫu theo thời gian, kênh, chi nhánh để tìm tình huống chung.", "Sau khi xác minh nguyên nhân, đặt người phụ trách, hạn hoàn thành và đo mức tái diễn."),
        }
        recommendation_cards = []
        for topic_id, data in topic_data["issues"][:3]:
            short, long = advice.get(data["recommendation_key"], advice["other"])
            title = escape(names[topic_id])
            quote = f"<br/><i>Trích dẫn: “{escape(data['sample'])}”</i>" if data.get("sample") else ""
            card = Table([
                [Paragraph(f"<b>{title}</b> ({data['negative']} phản hồi tiêu cực){quote}", body), ""],
                [Paragraph("<b>Ngắn hạn (7–14 ngày)</b><br/>" + escape(short), body), Paragraph("<b>Dài hạn</b><br/>" + escape(long), body)],
            ], colWidths=[252.5, 252.5])
            card.setStyle([
                ("BACKGROUND", (0, 0), (-1, -1), pale),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#d7e0e8")),
                ("INNERGRID", (0, 1), (-1, 1), 0.5, colors.HexColor("#d7e0e8")),
                ("SPAN", (0, 0), (1, 0)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
            ])
            recommendation_cards.extend([card, Spacer(1, 5)])
        content.extend(recommendation_cards or [Paragraph("Chưa có đủ phản hồi tiêu cực lặp lại để đưa ra khuyến nghị theo từng vấn đề.", body)])
        content.extend([
            Spacer(1, 8),
            Paragraph(
                "Tỷ lệ cảm xúc được tính trên nhận xét văn bản đã phân tích. Vấn đề lặp lại dựa trên chủ đề NLP và số phản hồi tiêu cực; cần đối chiếu trích dẫn với vận hành thực tế trước khi triển khai khuyến nghị. Báo cáo không đưa dữ liệu định danh khách hàng vào.",
                small,
            ),
        ])
        SimpleDocTemplate(
            output, pagesize=A4, leftMargin=44, rightMargin=44, topMargin=40, bottomMargin=40
        ).build(content)
    return output.getvalue()


async def run(tenant_id: str, job_id: str) -> None:
    resources = Resources.from_settings(get_settings())
    try:
        async with tenant_session(resources.session_factory, uuid.UUID(tenant_id)) as db:
            job = (
                await db.execute(
                    select(ExportJob)
                    .where(
                        ExportJob.id == uuid.UUID(job_id),
                        ExportJob.tenant_id == uuid.UUID(tenant_id),
                    )
                    .with_for_update(skip_locked=True)
                )
            ).scalar_one_or_none()
            if job is None or job.status not in (JobStatus.PENDING, JobStatus.RUNNING):
                return
            job.status = JobStatus.RUNNING
            try:
                if job.created_by is None:
                    raise RuntimeError("Người tạo tác vụ không còn tồn tại.")
                principal = await load_principal(db, job.created_by, job.tenant_id, None)
                from app.core.permissions import Permission

                principal.require(Permission.DATA_EXPORT)
                filters = FeedbackFilter.model_validate(job.params["filters"])
                content = await render(
                    AnalyticsService(db, principal), filters, job.kind, get_settings().app_name
                )
                suffix = "pdf" if job.kind == ExportKind.PDF else "xlsx"
                key = f"{tenant_id}/exports/{job_id}.{suffix}"
                await asyncio.to_thread(
                    LocalStorage(get_settings().storage_local_root).put, key, content
                )
                job.file_key = key
                job.filename = f"bao-cao-{job_id[:8]}.{suffix}"
                job.status = JobStatus.DONE
                job.finished_at = datetime.now(UTC)
            except Exception as exc:
                job.status = JobStatus.FAILED
                job.error_message = type(exc).__name__
    finally:
        await resources.close()


@celery_app.task(name="worker.tasks.exports.run")
def export_file(tenant_id: str, job_id: str) -> None:
    asyncio.run(run(tenant_id, job_id))


async def scan() -> int:
    resources = Resources.from_settings(get_settings())
    count = 0
    try:
        async with system_session(resources.session_factory) as db:
            tenants = list(
                (await db.execute(select(Tenant.id).where(Tenant.deleted_at.is_(None)))).scalars()
            )
        for tid in tenants:
            async with tenant_session(resources.session_factory, tid) as db:
                jobs = list(
                    (
                        await db.execute(
                            select(ExportJob.id).where(
                                ExportJob.tenant_id == tid,
                                ExportJob.status.in_([JobStatus.PENDING, JobStatus.RUNNING]),
                            )
                        )
                    ).scalars()
                )
            for jid in jobs:
                resources.queue.send(
                    "worker.tasks.exports.run", kwargs={"tenant_id": str(tid), "job_id": str(jid)}
                )
                count += 1
        return count
    finally:
        await resources.close()


@celery_app.task(name="worker.tasks.exports.scan")
def scan_exports() -> int:
    return asyncio.run(scan())
